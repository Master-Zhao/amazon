import csv
import io
from collections.abc import Iterator
from dataclasses import dataclass
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import BinaryIO

from openpyxl import load_workbook

from apps.advertising.models import EntityState, MatchType
from apps.reports.models import ReportType


CAMPAIGN_SCHEMA_VERSION = "campaign-v1"
TARGETING_SCHEMA_VERSION = "targeting-v1"
SEARCH_TERM_SCHEMA_VERSION = "search-term-v1"

COMMON_METRIC_ALIASES = {
    "date": ("date", "report_date", "report date"),
    "profile_id": ("profile_id", "profile id"),
    "campaign_id": ("campaign_id", "campaign id"),
    "campaign_name": ("campaign_name", "campaign name"),
    "currency": ("currency", "currency_code", "currency code"),
    "impressions": ("impressions",),
    "clicks": ("clicks",),
    "spend": ("spend", "cost"),
    "orders": ("orders", "purchases"),
    "sales": ("sales", "revenue"),
}

CAMPAIGN_COLUMN_ALIASES = {
    **COMMON_METRIC_ALIASES,
    "campaign_status": ("campaign_status", "campaign status", "state"),
    "daily_budget": ("daily_budget", "daily budget", "budget"),
    "snapshot_hour_local": (
        "snapshot_hour_local",
        "snapshot hour local",
        "report_hour_local",
    ),
}

TARGETING_COLUMN_ALIASES = {
    **COMMON_METRIC_ALIASES,
    "ad_group_id": ("ad_group_id", "ad group id"),
    "ad_group_name": ("ad_group_name", "ad group name"),
    "target_type": ("target_type", "target type"),
    "target_id": ("target_id", "target id", "keyword id"),
    "target_text": (
        "target_text",
        "target text",
        "keyword",
        "targeting expression",
    ),
    "match_type": ("match_type", "match type"),
    "target_status": ("target_status", "target status", "state"),
    "bid": ("bid", "target bid", "keyword bid"),
}

SEARCH_TERM_COLUMN_ALIASES = {
    **COMMON_METRIC_ALIASES,
    "ad_group_id": ("ad_group_id", "ad group id"),
    "ad_group_name": ("ad_group_name", "ad group name"),
    "search_term": ("search_term", "search term", "customer search term"),
    "targeting_expression": (
        "targeting_expression",
        "targeting expression",
        "keyword",
    ),
}

REQUIRED_CAMPAIGN_FIELDS = {
    "date",
    "campaign_id",
    "campaign_name",
    "campaign_status",
    "currency",
    "impressions",
    "clicks",
    "spend",
    "orders",
    "sales",
}

REQUIRED_TARGETING_FIELDS = {
    "date",
    "campaign_id",
    "campaign_name",
    "ad_group_id",
    "ad_group_name",
    "target_type",
    "target_id",
    "target_text",
    "target_status",
    "currency",
    "impressions",
    "clicks",
    "spend",
    "orders",
    "sales",
}

REQUIRED_SEARCH_TERM_FIELDS = {
    "date",
    "campaign_id",
    "campaign_name",
    "ad_group_id",
    "ad_group_name",
    "search_term",
    "targeting_expression",
    "currency",
    "impressions",
    "clicks",
    "spend",
    "orders",
    "sales",
}

REPORT_SCHEMAS = {
    ReportType.CAMPAIGN: (
        CAMPAIGN_SCHEMA_VERSION,
        CAMPAIGN_COLUMN_ALIASES,
        REQUIRED_CAMPAIGN_FIELDS,
    ),
    ReportType.TARGETING: (
        TARGETING_SCHEMA_VERSION,
        TARGETING_COLUMN_ALIASES,
        REQUIRED_TARGETING_FIELDS,
    ),
    ReportType.SEARCH_TERM: (
        SEARCH_TERM_SCHEMA_VERSION,
        SEARCH_TERM_COLUMN_ALIASES,
        REQUIRED_SEARCH_TERM_FIELDS,
    ),
}


def schema_version_for(report_type: str) -> str:
    try:
        return REPORT_SCHEMAS[report_type][0]
    except KeyError as exc:
        raise ReportStructureError(f"Unsupported report type: {report_type}") from exc


class ReportStructureError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class ParsedRow:
    row_number: int
    values: dict[str, str]


def _normalize_header(value: object) -> str:
    return str(value or "").strip().lower().replace("-", "_")


def _mapped_headers(
    headers: list[object],
    *,
    report_type: str,
) -> dict[str, int]:
    try:
        _, aliases_by_field, required_fields = REPORT_SCHEMAS[report_type]
    except KeyError as exc:
        raise ReportStructureError(f"Unsupported report type: {report_type}") from exc
    normalized = [_normalize_header(header) for header in headers]
    result: dict[str, int] = {}
    for canonical, aliases in aliases_by_field.items():
        for alias in aliases:
            if _normalize_header(alias) in normalized:
                result[canonical] = normalized.index(_normalize_header(alias))
                break
    missing = sorted(required_fields - result.keys())
    if missing:
        raise ReportStructureError(
            f"Missing required columns: {', '.join(missing)}"
        )
    return result


def _csv_rows(
    stream: BinaryIO,
    *,
    report_type: str,
) -> Iterator[ParsedRow]:
    wrapper = io.TextIOWrapper(stream, encoding="utf-8-sig", newline="")
    reader = csv.reader(wrapper)
    try:
        headers = next(reader)
    except StopIteration as exc:
        raise ReportStructureError("The report is empty") from exc
    mapped = _mapped_headers(headers, report_type=report_type)
    for row_number, row in enumerate(reader, start=2):
        if not any(str(value).strip() for value in row):
            continue
        values = {
            field: str(row[index] if index < len(row) else "").strip()
            for field, index in mapped.items()
        }
        yield ParsedRow(row_number=row_number, values=values)


def _xlsx_rows(
    stream: BinaryIO,
    *,
    report_type: str,
) -> Iterator[ParsedRow]:
    workbook = load_workbook(stream, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = list(next(rows))
        except StopIteration as exc:
            raise ReportStructureError("The workbook is empty") from exc
        mapped = _mapped_headers(headers, report_type=report_type)
        for row_number, row in enumerate(rows, start=2):
            if not any(str(value or "").strip() for value in row):
                continue
            values = {
                field: str(row[index] if index < len(row) and row[index] is not None else "").strip()
                for field, index in mapped.items()
            }
            yield ParsedRow(row_number=row_number, values=values)
    finally:
        workbook.close()


def iter_report_rows(
    *,
    stream: BinaryIO,
    filename: str,
    report_type: str,
) -> Iterator[ParsedRow]:
    if report_type not in REPORT_SCHEMAS:
        raise ReportStructureError(f"Unsupported report type: {report_type}")
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        yield from _csv_rows(stream, report_type=report_type)
        return
    if suffix == ".xlsx":
        yield from _xlsx_rows(stream, report_type=report_type)
        return
    raise ReportStructureError("Only .csv and .xlsx reports are accepted")


def _required(values: dict[str, str], field: str) -> str:
    value = values.get(field, "").strip()
    if not value:
        raise ValueError(f"{field} is required")
    return value


def _decimal(values: dict[str, str], field: str, *, required: bool = True) -> str | None:
    raw = values.get(field, "").strip()
    if not raw and not required:
        return None
    if not raw:
        raise ValueError(f"{field} is required")
    try:
        value = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"{field} must be a decimal") from exc
    if value < 0:
        raise ValueError(f"{field} must not be negative")
    return format(value, "f")


def _integer(values: dict[str, str], field: str) -> int:
    raw = _required(values, field)
    try:
        value = int(raw)
    except ValueError as exc:
        raise ValueError(f"{field} must be an integer") from exc
    if value < 0:
        raise ValueError(f"{field} must not be negative")
    return value


def normalize_campaign_row(
    row: ParsedRow,
    *,
    expected_profile_id: str,
    expected_currency: str,
) -> dict[str, object]:
    values = row.values
    profile_id = values.get("profile_id", "").strip()
    if profile_id and profile_id != expected_profile_id:
        raise ValueError("profile_id does not match the selected profile")
    try:
        report_date = date.fromisoformat(_required(values, "date"))
    except ValueError as exc:
        raise ValueError("date must use YYYY-MM-DD") from exc
    state = _required(values, "campaign_status").upper()
    if state not in EntityState.values:
        raise ValueError("campaign_status must be ENABLED, PAUSED, ARCHIVED or UNKNOWN")
    currency = _required(values, "currency").upper()
    if currency != expected_currency.upper():
        raise ValueError("currency does not match the selected profile")
    snapshot_hour_local = (
        _integer(values, "snapshot_hour_local")
        if values.get("snapshot_hour_local", "").strip()
        else None
    )
    if snapshot_hour_local is not None and snapshot_hour_local > 23:
        raise ValueError("snapshot_hour_local must be between 0 and 23")
    return {
        "date": report_date.isoformat(),
        "profile_id": expected_profile_id,
        "campaign_id": _required(values, "campaign_id"),
        "campaign_name": _required(values, "campaign_name"),
        "campaign_status": state,
        "daily_budget": _decimal(values, "daily_budget", required=False),
        "snapshot_hour_local": snapshot_hour_local,
        "currency": currency,
        "impressions": _integer(values, "impressions"),
        "clicks": _integer(values, "clicks"),
        "spend": _decimal(values, "spend"),
        "orders": _integer(values, "orders"),
        "sales": _decimal(values, "sales"),
    }


def _normalize_common_metric_row(
    values: dict[str, str],
    *,
    expected_profile_id: str,
    expected_currency: str,
) -> dict[str, object]:
    profile_id = values.get("profile_id", "").strip()
    if profile_id and profile_id != expected_profile_id:
        raise ValueError("profile_id does not match the selected profile")
    try:
        report_date = date.fromisoformat(_required(values, "date"))
    except ValueError as exc:
        raise ValueError("date must use YYYY-MM-DD") from exc
    currency = _required(values, "currency").upper()
    if currency != expected_currency.upper():
        raise ValueError("currency does not match the selected profile")
    return {
        "date": report_date.isoformat(),
        "profile_id": expected_profile_id,
        "campaign_id": _required(values, "campaign_id"),
        "campaign_name": _required(values, "campaign_name"),
        "currency": currency,
        "impressions": _integer(values, "impressions"),
        "clicks": _integer(values, "clicks"),
        "spend": _decimal(values, "spend"),
        "orders": _integer(values, "orders"),
        "sales": _decimal(values, "sales"),
    }


def normalize_targeting_row(
    row: ParsedRow,
    *,
    expected_profile_id: str,
    expected_currency: str,
) -> dict[str, object]:
    values = row.values
    normalized = _normalize_common_metric_row(
        values,
        expected_profile_id=expected_profile_id,
        expected_currency=expected_currency,
    )
    target_type = _required(values, "target_type").upper().replace(" ", "_")
    if target_type not in {"KEYWORD", "PRODUCT_TARGET"}:
        raise ValueError("target_type must be KEYWORD or PRODUCT_TARGET")
    state = _required(values, "target_status").upper()
    if state not in EntityState.values:
        raise ValueError("target_status must be ENABLED, PAUSED, ARCHIVED or UNKNOWN")
    match_type = values.get("match_type", "").strip().upper()
    if target_type == "KEYWORD" and match_type not in MatchType.values:
        raise ValueError("match_type must be BROAD, PHRASE or EXACT for KEYWORD")
    normalized.update(
        {
            "ad_group_id": _required(values, "ad_group_id"),
            "ad_group_name": _required(values, "ad_group_name"),
            "target_type": target_type,
            "target_id": _required(values, "target_id"),
            "target_text": _required(values, "target_text"),
            "match_type": match_type if target_type == "KEYWORD" else "",
            "target_status": state,
            "bid": _decimal(values, "bid", required=False),
        }
    )
    return normalized


def normalize_search_term_row(
    row: ParsedRow,
    *,
    expected_profile_id: str,
    expected_currency: str,
) -> dict[str, object]:
    values = row.values
    normalized = _normalize_common_metric_row(
        values,
        expected_profile_id=expected_profile_id,
        expected_currency=expected_currency,
    )
    normalized.update(
        {
            "ad_group_id": _required(values, "ad_group_id"),
            "ad_group_name": _required(values, "ad_group_name"),
            "search_term": _required(values, "search_term"),
            "targeting_expression": _required(values, "targeting_expression"),
        }
    )
    return normalized


def normalize_report_row(
    row: ParsedRow,
    *,
    report_type: str,
    expected_profile_id: str,
    expected_currency: str,
) -> dict[str, object]:
    normalizers = {
        ReportType.CAMPAIGN: normalize_campaign_row,
        ReportType.TARGETING: normalize_targeting_row,
        ReportType.SEARCH_TERM: normalize_search_term_row,
    }
    try:
        normalizer = normalizers[report_type]
    except KeyError as exc:
        raise ValueError(f"Unsupported report type: {report_type}") from exc
    return normalizer(
        row,
        expected_profile_id=expected_profile_id,
        expected_currency=expected_currency,
    )
