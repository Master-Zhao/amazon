from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"C:\QSZ\技术方案\_codex_plan\Amazon广告活动总览_远程数据联动_2小时验收实施计划.docx")

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
INK = "0F172A"
MUTED = "64748B"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
PALE = "F8FAFC"
WHITE = "FFFFFF"
GREEN = "166534"
AMBER = "92400E"
RED = "991B1B"
BORDER = "CBD5E1"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = 120) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_row_cant_split(row)
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def set_paragraph_keep(paragraph, *, keep_next=False, keep_lines=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if keep_next:
        keep = OxmlElement("w:keepNext")
        p_pr.append(keep)
    if keep_lines:
        keep = OxmlElement("w:keepLines")
        p_pr.append(keep)


def set_font(run, *, name="Calibri", east_asia="Microsoft YaHei", size=None,
             color=INK, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr()
    r_fonts = run._element.rPr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        run._element.rPr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), name)
    r_fonts.set(qn("w:hAnsi"), name)
    r_fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, fld_end])
    run2 = paragraph.add_run(" 页")
    set_font(run2, size=9, color=MUTED)


def set_doc_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def add_heading(doc: Document, text: str, level=1) -> None:
    p = doc.add_heading(text, level=level)
    set_paragraph_keep(p, keep_next=True)


def add_body(doc: Document, text: str, *, bold_label: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_label and text.startswith(bold_label):
        r1 = p.add_run(bold_label)
        set_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_label):])
        set_font(r2)
    else:
        r = p.add_run(text)
        set_font(r)


def add_callout(doc: Document, title: str, text: str, fill=LIGHT_BLUE, color=NAVY) -> None:
    table = doc.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title + "  ")
    set_font(r, bold=True, color=color)
    r = p.add_run(text)
    set_font(r, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_kv_table(doc: Document, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "项目"
    hdr[1].text = "结论"
    for cell in hdr:
        set_cell_shading(cell, LIGHT_BLUE)
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
    set_table_geometry(table, [2700, 6660])
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    set_font(run, size=9.5, bold=(r_idx == 0 or c_idx == 0))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_matrix(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], font_size=9.2) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = text
        set_cell_shading(cell, LIGHT_BLUE)
    for row_data in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cells[i].text = text
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        if r_idx % 2 == 0 and r_idx > 0:
            for cell in row.cells:
                set_cell_shading(cell, PALE)
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.08
                if c_idx == 0 and r_idx > 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_font(run, size=font_size, bold=(r_idx == 0), color=(NAVY if r_idx == 0 else INK))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_work_table(doc: Document, units: list[dict[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    headers = ["单元", "2 小时内实施内容与产物", "到点验收（必须留证）"]
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, BLUE)
    for unit in units:
        cells = table.add_row().cells
        cells[0].text = unit["id"] + "\n" + unit["time"]
        cells[1].text = "目标：" + unit["goal"] + "\n实施：" + unit["work"] + "\n产物：" + unit["output"]
        cells[2].text = unit["accept"]
    set_table_geometry(table, [820, 4100, 4440])
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(table.rows):
        if r_idx % 2 == 0 and r_idx > 0:
            for cell in row.cells:
                set_cell_shading(cell, PALE)
        for c_idx, cell in enumerate(row.cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    set_font(run, size=8.8, bold=(r_idx == 0 or (c_idx == 0 and r_idx > 0)), color=(WHITE if r_idx == 0 else INK))
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


WORK_PACKAGES = [
    (
        "工作包 A｜基线、并行任务接管与契约冻结（U01—U05，共 10 小时）",
        [
            {"id": "U01", "time": "第 1—2 小时", "goal": "建立页面验收基线。", "work": "对目标 HTML、用户截图与现有 Vue 页面逐项对照，固定模块、字段、交互和视觉参照；保存脱敏截图与差距清单，不读取隔离目录。", "output": "《页面基线与差距矩阵》：搜索栏、8 个 KPI、表现概览、风险卡、筛选标签、活动工具栏、19 列表格、合计、分页、弹窗/抽屉。", "accept": "矩阵每一项均有“HTML 现状 / 目标行为 / 数据来源 / 状态”；明确 HTML 当前无 fetch/axios、仅内置约 10 条数据并前端模拟日拆分。"},
            {"id": "U02", "time": "第 3—4 小时", "goal": "安全接管另一个对话的在途成果。", "work": "记录 Git dirty 文件、作者任务边界、已实现接口与未完成测试；在对方交付或明确移交前冻结同文件并行编辑。", "output": "交接清单与冲突所有权表，覆盖 backend advertising 四个文件、远程数据测试、AdvertisingOverviewPage 及新增 CampaignSection/API/types。", "accept": "逐文件记录 commit/diff 状态；现有约 964 行变更可追溯；没有覆盖、格式化或误删用户/另一任务改动。"},
            {"id": "U03", "time": "第 5—6 小时", "goal": "验证远程数据能支持页面。", "work": "只读采样 scm_remote.eb_ad_campaign 与 ads_analysis_remote.bi_analyze_ad_campaign，核对字段、空值、日期范围、商户映射、基数与索引；禁止输出账号、密钥和未脱敏行。", "output": "远程字段可用性与数据质量报告。", "accept": "至少验证 mer_id+mer_code 范围、campaign_id/code 回退、creation_date、impressions/click/spend/orders/sales/top-of-search 字段；异常只报告统计，不泄露真实明细。"},
            {"id": "U04", "time": "第 7—8 小时", "goal": "冻结字段与确定性公式。", "work": "把表格字段、KPI 和图表指标映射到远程列；固定金额 Decimal、业务日期、null 分母、币种与 Profile 时区语义。", "output": "字段/公式/格式化契约表。", "accept": "CTR=clicks/impressions、CPC=spend/clicks、CVR=orders/clicks、ACOS=spend/sales；分母无效返回 null/—；比例不得平均日百分比，金额不得跨币种相加。"},
            {"id": "U05", "time": "第 9—10 小时", "goal": "冻结 API 边界。", "work": "复核现有 campaigns 列表接口；推荐只新增一个 campaign-dashboard 聚合接口承载 KPI、每日趋势与风险摘要，导出复用同一筛选契约；写入 OpenAPI 草案。", "output": "接口决策记录、请求/响应样例和错误码。", "accept": "前端不直连远程库；所有请求带 tenant/profile 范围；列表与仪表盘使用同一日期/筛选；400/403/404/409/503 的展示行为明确。"},
        ],
    ),
    (
        "工作包 B｜远程 Campaign 列表纵向链路验收（U06—U10，共 10 小时）",
        [
            {"id": "U06", "time": "第 11—12 小时", "goal": "验收列表请求参数和权限。", "work": "完成/复核 DRF QuerySerializer：成对日期、90 个自然日上限、enabled/status/targetingType/search/ordering/page/pageSize/includeSummary；执行 Tenant、Store、Profile 授权。", "output": "稳定的请求校验层与边界测试。", "accept": "跨 Tenant/Profile 返回 404，范围内无权限返回 403；非法日期与排序字段在查询远程库前 400；pageSize 最大 100。"},
            {"id": "U07", "time": "第 13—14 小时", "goal": "验收远程聚合和元数据拼接。", "work": "按 mer_id+mer_code+campaign_id/code 聚合分析事实，用 SCM 补名称、状态、预算、竞价、开始/结束日期；远程库只读。", "output": "Selector/Adapter 聚合链路。", "accept": "无裸 campaign_id 越权查询；SCM 缺失时受控回退并标 metadataMatched/partialFields；远程异常统一为 503 且含 requestId。"},
            {"id": "U08", "time": "第 15—16 小时", "goal": "验收服务端搜索、枚举筛选和排序。", "work": "实现参数化名称/代码搜索，状态、启用、自动/手动投放筛选，以及白名单排序；验证 LIKE 通配符转义。", "output": "受控 SQL 筛选与排序测试。", "accept": "搜索 % 和 _ 不扩大范围；排序表达式只能来自白名单；同值使用 campaignKey 稳定排序；没有字符串拼接注入路径。"},
            {"id": "U09", "time": "第 17—18 小时", "goal": "验收分页、合计与元信息。", "work": "总数、当前页与筛选后 summary 分开查询；返回 source/currency/timezone/dataThroughDate/日期范围/语义标记。", "output": "CampaignListResponse 完整契约。", "accept": "分页合计针对全部筛选结果而非当前页；空结果 summary 为 null；totalPages/page 边界稳定；总成本明确为 spend 别名。"},
            {"id": "U10", "time": "第 19—20 小时", "goal": "验收现有 CampaignSection 前端基础。", "work": "接入 Axios 统一客户端、AbortController、URL 查询参数、loading/refreshing/empty/error/permission/partial 状态，核对 19 列与合计行。", "output": "可独立演示的远程活动表格。", "accept": "真实 API 返回渲染；刷新不闪空；后发请求胜出；空值显示 —；货币用 currency；前端测试通过且没有静态演示数组。"},
        ],
    ),
    (
        "工作包 C｜统一日期、搜索与条件筛选（U11—U15，共 10 小时）",
        [
            {"id": "U11", "time": "第 21—22 小时", "goal": "建立页面唯一查询状态。", "work": "将 tenant/store/marketplace/profile、日期、搜索、筛选、排序和页码收敛为单一查询模型；表格、KPI、图表、风险共用。", "output": "useCampaignOverviewQuery composable 或等价状态层。", "accept": "任一筛选只触发一轮协调刷新；URL 可复制恢复；切换 Profile 清空旧数据和选择；Pinia 不长期缓存业务事实。"},
            {"id": "U12", "time": "第 23—24 小时", "goal": "实现日期快捷与自定义筛选。", "work": "实现最近 7/14/30 天、自定义起止、包含首尾、90 天上限、无数据日期；按 Marketplace 本地业务日期提交。", "output": "可访问日期弹层与校验提示。", "accept": "开始晚于结束、缺一端、超上限均不发请求；选择后 KPI/图表/表格/风险日期一致；刷新 URL 后范围不丢失。"},
            {"id": "U13", "time": "第 25—26 小时", "goal": "完成表内广告活动搜索。", "work": "名称/代码搜索 300ms 防抖，Enter 立即提交，清空恢复，页码归 1；服务端查询，不在当前页本地过滤。", "output": "可验收的“查找广告活动”控件。", "accept": "快速输入只保留最后请求；中文、字母数字、空格和通配字符用例通过；搜索 chip 与 URL 同步。"},
            {"id": "U14", "time": "第 27—28 小时", "goal": "让顶部搜索栏成为真实交互。", "work": "实现 Ctrl/Cmd+K 聚焦和轻量命令/搜索面板：可定位 KPI、日期、筛选设置，并可提交 Campaign 名称/代码搜索；不伪装全站能力。", "output": "顶部搜索建议列表、键盘导航和无结果态。", "accept": "上下键/Enter/Escape 可用，焦点返回正确；提交 Campaign 搜索与表内搜索共用状态；未实现模块不返回假结果。"},
            {"id": "U15", "time": "第 29—30 小时", "goal": "完成条件筛选和筛选标签。", "work": "实现展示量、点击量、花费、订单、CPC、ACOS、CTR、CVR 的 >=、<=、>、<、=、between；后端参数化 HAVING/外层筛选；支持逐项删除与删除所有。", "output": "数值规则弹窗、动态 chips、OpenAPI 参数与测试。", "accept": "between 校验最小值<=最大值；比率输入/传输单位一致；所有条件作用于远程全量聚合而非当前页；删除后 KPI/图表/表格同步恢复。"},
        ],
    ),
    (
        "工作包 D｜顶部 KPI 仪表盘（U16—U20，共 10 小时）",
        [
            {"id": "U16", "time": "第 31—32 小时", "goal": "定义额外仪表盘接口。", "work": "完成 GET campaign-dashboard 契约：summary、trend、riskLevels、meta；请求复用列表日期和全部筛选，不接受任意商户参数。", "output": "OpenAPI、Serializer、TypeScript 生成类型。", "accept": "响应 camelCase；ID/金额/比例/null 语义固定；同一 Profile 单一币种；生成类型无未审查差异。"},
            {"id": "U17", "time": "第 33—34 小时", "goal": "实现 KPI 后端聚合。", "work": "从远程日事实汇总 impressions/clicks/orders/sales/spend，并由汇总值重算 CTR/CVR/ACOS/CPC；支持所有筛选。", "output": "仪表盘 summary Selector 与公式测试。", "accept": "结果与同筛选 Campaign 列表合计逐项相等；null 分母不返回 0/Infinity；Decimal 精度与币种正确。"},
            {"id": "U18", "time": "第 35—36 小时", "goal": "实现 8 个 KPI 卡片。", "work": "复刻展示量、点击量、购买量、销售额、花费、ACOS、CTR、CVR；补 tooltip、查看详情、骨架、空值与错误状态。", "output": "CampaignKpiBar 组件与单元测试。", "accept": "金额本地化且显示币种；比例保留约定小数；0 与无数据区分；没有任何 HTML 硬编码数字。"},
            {"id": "U19", "time": "第 37—38 小时", "goal": "处理趋势小图和对比语义。", "work": "仅当接口提供同长度前一周期数据时显示 delta/sparkline；否则隐藏比较，不沿用原型中的假上升箭头。", "output": "可选 comparison 契约与 KPI 趋势渲染。", "accept": "前期区间按业务日期正确推导且不越数据范围；正负/无变化样式正确；无比较数据时页面不暗示涨跌。"},
            {"id": "U20", "time": "第 39—40 小时", "goal": "验证 KPI 全量联动。", "work": "覆盖日期、搜索、状态、启用、类型、数值条件和 Profile 切换，验证请求竞态与局部失败。", "output": "KPI 联动集成测试和证据截图。", "accept": "任何条件变化后的 KPI 与表格合计一致；旧请求不能覆盖新范围；503 显示可重试，不保留旧日期的 KPI 冒充新结果。"},
        ],
    ),
    (
        "工作包 E｜中部表现概览趋势图（U21—U25，共 10 小时）",
        [
            {"id": "U21", "time": "第 41—42 小时", "goal": "实现每日趋势聚合。", "work": "按 Profile 业务日期生成连续日期桶；绝对值求和，比率由每日日原始分子分母重算；缺失日定义为 0 或 null 并记录语义。", "output": "daily trend Selector 与边界测试。", "accept": "包含起止日；跨月、闰日、无数据日正确；绝不把总量均摊到每天；日比例不做平均。"},
            {"id": "U22", "time": "第 43—44 小时", "goal": "完成趋势接口序列化和性能界限。", "work": "返回 date、impressions、clicks、orders、sales、spend、ctr/cvr/acos/cpc；复用权限与筛选；限制 90 个点。", "output": "dashboard trend API 与查询次数/耗时基线。", "accept": "固定数量 SQL、无 N+1；空区间返回空数组；取消请求安全；敏感远程映射不出现在响应。"},
            {"id": "U23", "time": "第 45—46 小时", "goal": "搭建 ECharts 表现概览。", "work": "创建可销毁/可 resize 的图表组件，默认 4 个指标，展示日期范围、图例、双轴与 tooltip。", "output": "CampaignPerformanceChart 基础组件。", "accept": "首次加载、容器缩放、路由离开均无控制台错误或实例泄漏；空数据/加载/失败态不叠图。"},
            {"id": "U24", "time": "第 47—48 小时", "goal": "实现四槽指标选择。", "work": "8 个指标可切换到 4 个槽位；金额、计数和比例使用匹配轴/格式；颜色与图例稳定，避免重复选择歧义。", "output": "指标选择器、系列配置与测试。", "accept": "切换不重新请求数据；tooltip 数字单位正确；左右轴不会把百分比与金额混为一轴；图例总值与 KPI 一致。"},
            {"id": "U25", "time": "第 49—50 小时", "goal": "完成图表交互与可访问性。", "work": "实现显示全部/仅图表/全部隐藏、放大视图、键盘可操作说明和数据表替代视图。", "output": "完整表现概览交互。", "accept": "三种视图切换可恢复；放大不改变数据；键盘和读屏可获取指标/日期/值；小屏不截断关键控件。"},
        ],
    ),
    (
        "工作包 F｜广告活动表格完整交互（U26—U30，共 10 小时）",
        [
            {"id": "U26", "time": "第 51—52 小时", "goal": "完成字段与提示文案。", "work": "对齐选择框、启用、名称/代码、投放类型、状态、竞价、起止、日预算、展示、首页位置、花费、点击、CTR、总成本、订单、CPC、ACOS、CVR；表头 tooltip 解释口径。", "output": "字段齐全的横向滚动表格。", "accept": "19 列顺序和截图/HTML 对齐；短值不换行，名称可省略并可查看全文；表头与合计行滚动对齐。"},
            {"id": "U27", "time": "第 53—54 小时", "goal": "完成排序、分页和行选择。", "work": "白名单表头排序、15/30/50 页大小、上一页/下一页/页码、当前页全选与跨页选择说明。", "output": "稳定表格状态与选择状态测试。", "accept": "排序/筛选后页码归 1；全选只选择当前页或明确提示跨页策略；选择不改变 KPI；无越界页。"},
            {"id": "U28", "time": "第 55—56 小时", "goal": "实现行级详情入口。", "work": "名称/展开图标进入 Campaign 详情或侧栏；服务端 campaignKey 再次绑定 tenant/profile/mer_id/mer_code；显示部分字段来源状态。", "output": "安全的详情导航/抽屉骨架。", "accept": "篡改 campaignKey 返回 404；返回页面保留查询条件与滚动位置；无 SCM 元数据时仍能展示分析数据。"},
            {"id": "U29", "time": "第 57—58 小时", "goal": "正确处理创建和启停按钮。", "work": "依据 V1 禁止真实 Amazon Ads 写入：创建按钮和开关不得硬编码成功；采用禁用+原因，或在已有 Action Preview/人工执行链路授权时跳转。", "output": "明确的只读/动作预览交互。", "accept": "点击不会写远程库或伪造状态；权限不足按钮隐藏/禁用但后端仍校验；页面清楚标识只读数据源。"},
            {"id": "U30", "time": "第 59—60 小时", "goal": "实现按当前条件导出。", "work": "后端流式导出筛选后全量字段，前端携带同一日期/搜索/筛选/排序；限制规模并记录 requestId。", "output": "CSV 导出接口与下载交互。", "accept": "导出行数/合计与页面筛选一致；CSV UTF-8 BOM/表头/金额/日期规则明确；公式注入字符被安全处理；超限给出可操作错误。"},
        ],
    ),
    (
        "工作包 G｜指标详情、风险评估与异常联动（U31—U35，共 10 小时）",
        [
            {"id": "U31", "time": "第 61—62 小时", "goal": "实现 KPI 查看详情抽屉。", "work": "复用 trend 数据展示选定指标趋势、总值和 Campaign 排名；7 个/8 个 KPI 共用一个配置化抽屉。", "output": "MetricDetailDrawer 与配置表。", "accept": "标题、图标、单位、排名字段随指标正确变化；日期/筛选与主页面一致；打开抽屉不额外重复拉取可复用数据。"},
            {"id": "U32", "time": "第 63—64 小时", "goal": "冻结风险规则与数据门槛。", "work": "只使用确定性规则和已确认阈值，定义 LOW/MEDIUM/HIGH、INSUFFICIENT_DATA、证据字段和配置版本；AI 只解释。", "output": "风险规则契约、阈值来源和版本化测试夹具。", "accept": "无足够点击/花费/订单时不判正常或异常；每项风险能追溯到 Campaign、日期、公式、阈值和配置版本。"},
            {"id": "U33", "time": "第 65—66 小时", "goal": "实现风险汇总与异常详情数据。", "work": "在 campaign-dashboard 返回风险等级计数/占比和有限异常摘要；详情按需查询且仍受 Profile 范围。", "output": "风险 Selector/API 与隔离测试。", "accept": "风险数等于筛选范围内规则结果；无异常显示 0 而非静态示例；跨 Tenant 404；查询受上限和稳定排序保护。"},
            {"id": "U34", "time": "第 67—68 小时", "goal": "复刻风险卡、异常弹窗和归因。", "work": "实现等级条、查看全部、Campaign 跳转、归因证据列表；修复建议仅连接现有 Recommendation/Action Preview 能力。", "output": "RiskAssessmentCard、RiskModal、AttributionModal。", "accept": "等级颜色与文案一致；点击风险可定位对应活动；没有 LLM 隐藏思维；无建议能力时明确不可用而非 alert 演示。"},
            {"id": "U35", "time": "第 69—70 小时", "goal": "验证页面各模块一致性和部分失败。", "work": "模拟列表成功/仪表盘失败、仪表盘成功/风险失败、远程超时和 Profile 切换，定义局部重试。", "output": "页面状态机与故障矩阵。", "accept": "一个模块失败不抹掉其他已确认模块；旧范围数据不会在新范围下继续显示；每个错误显示 requestId 与重试入口。"},
        ],
    ),
    (
        "工作包 H｜自动化验收、性能、安全与交付（U36—U40，共 10 小时）",
        [
            {"id": "U36", "time": "第 71—72 小时", "goal": "补齐前端单元与组件测试。", "work": "覆盖日期、搜索防抖、条件筛选、chips、KPI、图表选择、表格排序分页、抽屉/弹窗、空/错/权限/竞态。", "output": "Vitest 测试集。", "accept": "pnpm --dir frontend test、typecheck、lint 通过；每类核心交互至少有成功与失败断言；无依赖真实远程库的前端测试。"},
            {"id": "U37", "time": "第 73—74 小时", "goal": "补齐后端 API 与隔离测试。", "work": "覆盖参数、公式、null、币种、日期、搜索转义、数值 HAVING、排序、分页、summary/trend/risk/export、远程故障与权限。", "output": "pytest 契约/安全/数据语义测试。", "accept": "目标测试通过；同时断言远程数据库路由禁止迁移/写入；SQL 注入与跨租户用例通过。"},
            {"id": "U38", "time": "第 75—76 小时", "goal": "完成契约、构建和查询性能门槛。", "work": "生成 OpenAPI/TS 类型并检查差异；执行 Django check、迁移检查、前端 build；对典型 30/90 天查询记录 EXPLAIN 与上限。", "output": "自动化命令日志、类型差异和性能证据。", "accept": "标准检查全绿；固定查询数、无 N+1/无界扫描；未验证项明确标 NOT VERIFIED，不虚构毫秒或 QPS。"},
            {"id": "U39", "time": "第 77—78 小时", "goal": "完成真实浏览器视觉与交互验收。", "work": "在 1656×707 参照尺寸及常见桌面宽度执行 Playwright/浏览器验收，检查横向滚动、粘性表头、弹层、键盘、刷新恢复和截图对比。", "output": "E2E 脚本、通过/失败统计与视觉证据。", "accept": "截图核心布局一致；无重叠、裁切、控制台错误；日期→KPI→图表→表格→风险完整链路真实走通。"},
            {"id": "U40", "time": "第 79—80 小时", "goal": "完成文档、演示和交付收口。", "work": "更新 OpenAPI、决策日志、页面字段口径、远程数据说明、测试报告、启动方法、演示脚本和遗留风险；与另一任务最终 diff 复核。", "output": "可交接发布包和阶段报告。", "accept": "按 AGENTS.md 17 项报告格式交付；代码/测试/文档状态一致；所有未执行项有原因；不提前宣称真实 Amazon 写操作已完成。"},
        ],
    ),
]


def build() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    set_doc_styles(doc)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("AMAZON 广告活动总览  |  实施计划")
    set_font(hr, size=9, bold=True, color=MUTED)
    footer = section.footer
    fp = footer.paragraphs[0]
    add_page_number(fp)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("AMAZON 广告活动总览")
    set_font(r, size=26, color=NAVY, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("远程数据联动与完整交互 · 2 小时验收实施计划")
    set_font(r, size=15, color=BLUE, bold=True)

    add_kv_table(doc, [
        ("计划总量", "40 个验收单元 × 2 小时 = 80 小时净实施时间；每个单元到点必须可独立演示、测试或审阅。"),
        ("输入基线", r"C:\Users\admin\Desktop\广告智能投放系统\广告智能投放系统\index.html 与用户截图。"),
        ("工程基线", r"C:\QSZ\技术方案；遵守主规格、PLANS.md、docs/15-decision-log.md 与 AGENTS.md。"),
        ("数据原则", "浏览器只调用 Django /api/v1；Django 使用只读远程数据库别名查询；远程库不得迁移、写入或直接暴露。"),
        ("编制日期", "2026-08-04（Asia/Shanghai）"),
    ])

    add_callout(
        doc,
        "关键结论",
        "目标 HTML 是 2,416 行的静态原型：没有远程请求，约 10 条 Campaign 内嵌在脚本中，日期趋势由前端模拟拆分。现有工程中另一个任务已开始实现远程 Campaign 列表接口和表格组件；本计划以先验收/接管该成果为门槛，再补齐 KPI、图表、全局搜索、数值筛选、风险、详情和导出。",
    )

    add_heading(doc, "1. 范围与强制边界", 1)
    add_body(doc, "范围：复刻目标页面中可见的搜索、KPI、趋势图、风险评估、筛选标签、活动工具栏、日期选择、条件筛选、字段表格、合计、分页、详情弹窗/抽屉，并让全部业务数据来自远程数据库经后端 API 的受控读取。", bold_label="范围：")
    add_body(doc, "真实写操作边界：V1 不调用 Amazon Ads API，也不允许改写现有远程库。因此“创建广告活动”和“启用/暂停”只能做权限化禁用说明，或跳转到已经存在的 Action Preview / 人工执行链路；绝不能弹出硬编码成功或即时改变远程状态。", bold_label="真实写操作边界：")
    add_body(doc, "一致性边界：KPI、图表、风险和列表必须使用同一 Tenant → Store → Marketplace → AdvertisingProfile、同一日期范围和同一筛选集合。金额不跨币种汇总，业务日期使用 Marketplace 语义。", bold_label="一致性边界：")
    add_body(doc, "并行协作边界：另一个任务当前占用 advertising selectors/serializers/views、remote_databases、remote advertising tests 以及 AdvertisingOverviewPage、CampaignSection/API/types。U02 未完成前禁止多人同时编辑这些文件。", bold_label="并行协作边界：")

    add_heading(doc, "2. 已核对的 HTML 功能现状", 1)
    add_matrix(doc, ["模块", "原型现状", "落地要求"], [
        ["顶部搜索", "仅 console.log；Ctrl/Cmd+K 能聚焦", "真实搜索/命令面板；与 Campaign 服务端搜索共享状态"],
        ["8 个 KPI", "由 generateMockData 计算；含假 delta/小趋势", "远程聚合；null/币种/比较周期语义真实"],
        ["表现概览", "原生 SVG；4 槽可切 8 指标", "ECharts；每日真实趋势、双轴、tooltip、resize"],
        ["风险评估", "静态 RISK_EXCEPTIONS/ATTRIBUTION_DATA", "确定性规则、数据门槛、版本与证据可追溯"],
        ["日期", "7/14/30 天和自定义；前端模拟刷新", "服务端 inclusive 日期；90 日上限；Profile 业务日期"],
        ["条件筛选", "数值规则本地过滤；状态分支仍为占位", "远程全量聚合后服务端筛选；chips 全联动"],
        ["活动表", "约 10 条内嵌数据；前端分页", "远程服务端搜索、筛选、排序、分页、合计"],
        ["创建/启停/导出", "创建 alert；开关只改内存；导出无事件", "写操作不伪造；导出真实筛选结果"],
        ["指标详情", "通用抽屉可演示", "复用远程趋势与 Campaign 排名"],
    ], [1700, 3400, 4260], font_size=9.0)

    add_heading(doc, "3. 推荐数据流与接口分工", 1)
    add_callout(doc, "推荐最小接口组合", "保留正在实现的 Campaign 列表接口，再新增 1 个 campaign-dashboard 聚合接口；导出使用独立下载路由但复用相同筛选 Serializer。这样不会让分页表格承担每日趋势，也避免前端把远程事实重新汇总。", fill=LIGHT_GRAY)
    add_matrix(doc, ["接口/数据流", "职责", "关键要求"], [
        ["GET …/campaigns", "表格 items、筛选后 summary、分页、meta", "服务端搜索/筛选/排序；summary 是全部筛选结果，不是当前页"],
        ["GET …/campaign-dashboard", "8 KPI、每日 trend、riskLevels/有限异常摘要、meta", "与 campaigns 共用查询条件；比率由原始分子分母重算"],
        ["GET …/campaigns/export", "当前筛选的全量 CSV", "流式、规模上限、公式注入防护、UTF-8 规则"],
        ["前端统一查询状态", "协调两个读取接口、URL、取消与局部重试", "后发请求胜出；切 Profile 不显示旧数据"],
        ["远程数据库", "scm_remote 元数据 + ads_analysis_remote 日事实", "SELECT only；mer_id+mer_code 强制范围；无前端直连"],
    ], [2500, 3300, 3560], font_size=8.9)

    add_heading(doc, "4. 字段、KPI 与公式验收口径", 1)
    add_matrix(doc, ["页面字段", "来源/公式", "显示与空值"], [
        ["启用/状态", "SCM state 优先，分析状态回退并规范化", "只读开关；状态标签；未知值不猜测"],
        ["活动名称/代码", "SCM name/code，分析 campaign_name/code 回退", "名称主行、代码辅行；可搜索"],
        ["投放类型/竞价", "SCM type/bidding_strategy，分析字段回退", "AUTO/MANUAL 与竞价枚举中文化"],
        ["开始/结束日期", "SCM 当前元数据；分析创建日回退", "Profile 业务日期；无结束日期单独文案"],
        ["日预算", "SCM 当前预算优先，分析最新快照回退", "Money{amount,currencyCode}；不得跨日求和"],
        ["展示量", "SUM(impressions/远程实际列)", "整数千分位；有数据的 0 显示 0"],
        ["点击量", "SUM(clicks)", "整数千分位"],
        ["花费/总成本", "SUM(spend)；总成本为 spend 别名", "货币格式；契约显式标 SPEND_ALIAS"],
        ["购买量", "SUM(orders)", "整数；需标注远程归因字段语义"],
        ["销售额", "SUM(sales)", "KPI/图表/详情使用同一币种"],
        ["首页位置", "远程 top-of-search 字段按已确认口径聚合", "口径未验证时标记 REMOTE_FIELDS_UNVERIFIED"],
        ["CTR", "clicks ÷ impressions", "分母 0 → null/—；百分比"],
        ["CPC", "spend ÷ clicks", "分母 0 → null/—；货币"],
        ["CVR", "orders ÷ clicks", "分母 0 → null/—；百分比"],
        ["ACOS", "spend ÷ sales", "分母 0 → null/—；百分比"],
    ], [1900, 4200, 3260], font_size=8.7)

    add_heading(doc, "5. 80 小时逐单元实施与验收计划", 1)
    add_body(doc, "执行规则：每个单元严格 2 小时。到点若未满足验收，不把未完成内容悄悄带入下一单元；记录失败证据、缩小范围或重新排期。U01—U10 是在途成果的接管与纵向链路门槛，后续单元不得绕过。")
    for title, units in WORK_PACKAGES:
        doc.add_page_break()
        add_heading(doc, title, 2)
        add_work_table(doc, units)

    doc.add_page_break()
    add_heading(doc, "6. 每个 2 小时单元的统一验收证据", 1)
    add_matrix(doc, ["证据", "最低要求"], [
        ["变更边界", "列出实际修改/新增文件；确认未进入 amazon-ads-operations-0.1.1；说明是否与并行任务重叠。"],
        ["可运行结果", "给出可访问页面/API，真实返回成功、空数据和失败至少一类；不得使用静态成功响应。"],
        ["自动化", "记录实际命令、退出码、通过/失败/跳过数量；未执行必须写原因。"],
        ["契约", "请求参数、响应字段、null/金额/币种/日期/分页语义与 OpenAPI 一致。"],
        ["安全", "Tenant/Store/Profile 范围、远程 SELECT-only、无秘密/真实明细泄露、无字符串拼接 SQL。"],
        ["视觉/交互", "页面状态、键盘、焦点、横向滚动、弹层、错误和 requestId 实际可见。"],
        ["交接", "下一单元依赖、遗留风险和需要人工确认的决策不超过三条，且有明确负责人。"],
    ], [1900, 7460], font_size=9.2)

    add_heading(doc, "7. 阶段性放行门槛", 1)
    add_matrix(doc, ["里程碑", "完成单元", "放行条件"], [
        ["M1 在途成果可接管", "U01—U10 / 20h", "远程列表纵向链路、权限、分页、合计、基础表格测试可重复；并行编辑冲突解除。"],
        ["M2 查询联动可用", "U11—U15 / 30h", "日期、顶部/表内搜索、状态/类型/启用/数值条件统一作用于远程全量。"],
        ["M3 仪表盘可用", "U16—U25 / 50h", "8 KPI 与每日趋势真实、同筛选一致、无模拟拆日或假涨跌。"],
        ["M4 页面完整", "U26—U35 / 70h", "表格字段/导出/详情/风险/抽屉和局部失败处理闭环。"],
        ["M5 可交付", "U36—U40 / 80h", "测试、类型、构建、OpenAPI、性能/安全、浏览器与文档证据齐全。"],
    ], [2200, 1900, 5260], font_size=9.0)

    doc.add_page_break()
    add_heading(doc, "8. 必须提前确认或持续标注的风险", 1)
    add_matrix(doc, ["风险/决策", "处理原则", "阻断点"], [
        ["另一个任务尚未交付", "只审阅，不并行覆盖；U02 建立交接点", "U06 前"],
        ["远程字段拼写/空值/归因未知", "U03 用只读统计验证；未验证字段在 UI/meta 显式标识", "相关 KPI/风险放行前"],
        ["top-of-search 聚合口径", "不得猜 MAX/AVG/SUM；以真实字段语义和样例确认", "U04/U17"],
        ["风险阈值", "必须来自已确认配置/规则版本；数据不足返回 INSUFFICIENT_DATA", "U32"],
        ["创建/启停", "V1 不写 Amazon/远程库；只允许禁用说明或 Action Preview/人工执行", "U29"],
        ["导出规模", "先设置可观测上限和流式策略，不做无界全库导出", "U30"],
        ["不同币种/Marketplace", "按 Profile 分开显示，不直接相加", "全程"],
        ["真实性能", "记录真实 EXPLAIN/耗时，未跑压测不得写 QPS 或延迟结论", "U38"],
    ], [2500, 5200, 1660], font_size=9.0)

    doc.add_page_break()
    add_heading(doc, "9. 最终标准检查命令", 1)
    commands = [
        "uv sync --project backend --frozen",
        "uv run --project backend python backend/manage.py check",
        "uv run --project backend python backend/manage.py makemigrations --check --dry-run",
        "uv run --project backend python backend/manage.py migrate --check",
        "uv run --project backend pytest backend",
        "pnpm --dir frontend install --frozen-lockfile",
        "pnpm --dir frontend generate:api",
        "pnpm --dir frontend lint",
        "pnpm --dir frontend typecheck",
        "pnpm --dir frontend test",
        "pnpm --dir frontend build",
        "pnpm --dir frontend test:e2e",
        "docker compose -f compose.local.yml config --quiet",
        "docker compose -f compose.test.yml config --quiet",
        "docker compose --env-file .env.example -f compose.prod.yml config --quiet",
    ]
    add_matrix(doc, ["序号", "命令"], [[str(i + 1), cmd] for i, cmd in enumerate(commands)], [720, 8640], font_size=9.0)
    add_callout(doc, "交付纪律", "以上命令只有实际执行并记录真实结果后才能写“通过”。若远程环境、凭据或服务不可用，应写明 NOT VERIFIED 与原因，不得以 mock 或原型页面替代远程数据库联调证据。", fill="FEF3C7", color=AMBER)

    doc.core_properties.title = "Amazon 广告活动总览远程数据联动 2 小时验收实施计划"
    doc.core_properties.subject = "基于指定 HTML 与截图的 80 小时实施与验收拆解"
    doc.core_properties.author = "Codex"
    doc.core_properties.keywords = "Amazon Ads, Campaign, Remote MySQL, Dashboard, KPI, ECharts, 2-hour acceptance"
    doc.core_properties.comments = "compact_reference_guide preset; CJK font override: Microsoft YaHei; schedule table compact override: 8.8pt."
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
