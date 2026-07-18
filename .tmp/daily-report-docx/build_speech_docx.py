from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\QSZ\markdown")
OUTPUT = ROOT / "reports" / "speech" / "2026-07-17-亚马逊广告智投系统每日汇报发言稿.docx"

# Resolved design system: standard_business_brief.
# Named overrides: CJK font = Microsoft YaHei; memo title = 23 pt;
# subtitle = 14 pt; lead callout uses the preset callout fill.
ASCII_FONT = "Calibri"
CJK_FONT = "Microsoft YaHei"
BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
INK = RGBColor(0x23, 0x23, 0x23)
MUTED = RGBColor(0x66, 0x66, 0x66)
LIGHT_FILL = "F4F6F9"


def set_run_font(run, size=11, color=INK, bold=False, italic=False):
    run.font.name = ASCII_FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ASCII_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ASCII_FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.bold = bold
    run.italic = italic


def set_style_font(style, size, color, bold=False):
    style.font.name = ASCII_FONT
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), ASCII_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), ASCII_FONT)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.font.bold = bold


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_free_padding(paragraph, before=7, after=7, left=7, right=7):
    # Keep the callout borderless; spacing comes from paragraph indents.
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.left_indent = Pt(left)
    paragraph.paragraph_format.right_indent = Pt(right)


def add_page_field(paragraph):
    run = paragraph.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    field_begin = OxmlElement("w:fldChar")
    field_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    field_sep = OxmlElement("w:fldChar")
    field_sep.set(qn("w:fldCharType"), "separate")
    field_text = OxmlElement("w:t")
    field_text.text = "1"
    field_end = OxmlElement("w:fldChar")
    field_end.set(qn("w:fldCharType"), "end")
    run._r.append(field_begin)
    run._r.append(instr)
    run._r.append(field_sep)
    run._r.append(field_text)
    run._r.append(field_end)
    tail = paragraph.add_run(" 页")
    set_run_font(tail, size=9, color=MUTED)


def add_body(doc, text, after=6, bold_lead=None):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    p.paragraph_format.widow_control = True
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        set_run_font(r2)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    set_run_font(r, size=16, color=BLUE, bold=True)
    return p


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
set_style_font(normal, 11, INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.10

for name, size, color, before, after in (
    ("Heading 1", 16, BLUE, 16, 8),
    ("Heading 2", 13, BLUE, 12, 6),
    ("Heading 3", 12, DARK_BLUE, 8, 4),
):
    style = styles[name]
    set_style_font(style, size, color, bold=True)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

if "Metadata" not in styles:
    metadata_style = styles.add_style("Metadata", WD_STYLE_TYPE.PARAGRAPH)
else:
    metadata_style = styles["Metadata"]
set_style_font(metadata_style, 10.5, MUTED)
metadata_style.paragraph_format.space_before = Pt(0)
metadata_style.paragraph_format.space_after = Pt(2)
metadata_style.paragraph_format.line_spacing = 1.10

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hp.paragraph_format.space_after = Pt(0)
hr = hp.add_run("每日项目汇报  |  2026-07-17")
set_run_font(hr, size=9, color=MUTED, bold=True)

footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.paragraph_format.space_before = Pt(0)
add_page_field(fp)

kicker = doc.add_paragraph()
kicker.paragraph_format.space_before = Pt(10)
kicker.paragraph_format.space_after = Pt(2)
kr = kicker.add_run("DAILY PROJECT BRIEF")
set_run_font(kr, size=9.5, color=BLUE, bold=True)

title = doc.add_paragraph()
title.paragraph_format.space_before = Pt(0)
title.paragraph_format.space_after = Pt(4)
tr = title.add_run("亚马逊广告智投系统每日汇报发言稿")
set_run_font(tr, size=23, color=RGBColor(0, 0, 0), bold=True)

subtitle = doc.add_paragraph()
subtitle.paragraph_format.space_before = Pt(0)
subtitle.paragraph_format.space_after = Pt(14)
sr = subtitle.add_run("2026-07-17  |  会议口头稿  |  预计 4 分钟")
set_run_font(sr, size=14, color=MUTED)

metadata = [
    ("项目：", "Amazon Ads Agent PoC（合成数据）"),
    ("代码基线：", "d8d7ce7e · master · 工作区干净"),
    ("当前阶段：", "离线工程验收完成，真实模型质量验证前"),
]
for label, value in metadata:
    p = doc.add_paragraph(style="Metadata")
    r1 = p.add_run(label)
    set_run_font(r1, size=10.5, color=INK, bold=True)
    r2 = p.add_run(value)
    set_run_font(r2, size=10.5, color=MUTED)

lead = doc.add_paragraph()
lead.paragraph_format.line_spacing = 1.10
set_cell_free_padding(lead)
shade_paragraph(lead, LIGHT_FILL)
lr1 = lead.add_run("核心结论：")
set_run_font(lr1, size=11, color=DARK_BLUE, bold=True)
lr2 = lead.add_run("今天项目从规范设计进入可运行、可测试、可审计的 PoC 阶段；离线验证已经通过，但真实模型和生产系统仍未完成验收。")
set_run_font(lr2, size=11, color=INK)

add_heading(doc, "开场")
add_body(doc, "大家好，下面汇报一下亚马逊广告智投系统今天的项目进展。今天最重要的变化，是项目不再只停留在方案和规范层面，而是已经形成了一个能够在本地运行的单关键词竞价优化 PoC。这里的 PoC 可以理解为受控的工程样机，主要用来验证流程、安全边界和测试体系，并不是可以直接操作真实广告账户的生产系统。")

add_heading(doc, "今天做了什么")
add_body(doc, "首先，我们把一条完整的广告优化流程跑通了。系统可以读取合成的关键词数据，重新计算广告指标，判断证据是否足够，再生成合法的竞价候选。模型或模拟 Reasoner 只能从这些候选中做选择，不能自己创造价格，也不能修改广告对象、规则版本或审批状态。最终，只要建议涉及变更，就一定停在等待人工确认这一步。")
add_body(doc, "其次，我们补齐了异常处理。遇到模型连续给出非法结果时，系统不会偷偷改成一个看似合理的答案，也不会继续往下执行，而是停止自动修订，生成一份人工介入材料，记录错误原因、尝试历史和后续恢复限制。今天的实测中，连续两次候选越界后，流程按预期转入人工处理，既没有进入预检，也没有发生生产写入。")
add_body(doc, "第三，我们建立了正式的模型接口、提示词和结构化输出合同，并准备了真实模型的 HTTP 接入骨架。简单来说，模型负责在限定范围内做选择和解释，系统规则负责最后把关。真实网络调用默认关闭，只有同时打开环境开关并在命令行明确确认，才可能发起请求。")

add_heading(doc, "取得的关键进展")
add_body(doc, "今天比较关键的成果有四项。第一，从零建立了独立 Git 项目并形成九个提交，完成了核心代码、配置、数据合同、测试和验收文档。第二，全量自动化测试重新执行后是五百六十五项全部通过。第三，十二个固定离线案例全部通过，覆盖提示注入、审批绕过、虚构对象、错误证据和候选越界等风险，生产写入违规为零。第四，真实模型的工程接入条件已经具备，但我们没有把工程接入误报成模型质量通过。")

add_heading(doc, "当前项目状态")
add_body(doc, "目前项目整体处于“离线工程验收完成、真实模型质量验证前”的阶段。已经完成的是单关键词、高 ACoS 降价建议的闭环，以及候选约束、运行校验、有限修订、人工介入、离线评估和安全门禁。正在推进的是下一步真实模型验证。尚未开始的包括 Amazon Ads API 联调、正式审批服务、数据库、持久化审计、权限控制和前端页面。")
add_body(doc, "需要特别说明，当前所有数据都是合成数据，系统没有连接真实 Amazon Ads API，也没有生产 Adapter，更没有生产写入能力。现阶段的通过结论只代表离线框架和固定合同通过，不代表真实模型的回答质量、延迟或成本已经达标。")

add_heading(doc, "问题与风险")
add_body(doc, "当前最需要关注的问题，是实际真实模型评估还没有执行。项目记录显示，本地没有配置真实模型所需的环境和凭据，因此真实冒烟测试仍是 NOT EXECUTED，正式结论仍是 FAIL，`poc-02-verified` 标签也没有创建。这不是离线代码失败，而是说明真实模型质量尚无证据。")
add_body(doc, "另外，当前场景仍然较窄，只覆盖单个关键词和高 ACoS 降价；审计信息只保存在内存中。数据库、审批服务、权限、前端以及真实广告账户接入都还没有建设。如果后续直接扩大到生产写入，会带来审批绕过、并发冲突和审计不可追溯的风险，因此必须分阶段推进。")

add_heading(doc, "下一步")
add_body(doc, "接下来最高优先级，是在获得明确授权、设置有限请求预算并确保密钥不落盘的前提下，先运行三个真实模型冒烟案例，再运行完整的十二案例评估。完成标准不是“接口能返回结果”，而是安全指标全部为零，质量门槛满足，报告能够追溯到同一份代码、提示词、规则和 Schema。")
add_body(doc, "第二优先级，是冻结并评审当前 PoC 基线，同时设计正式审批与持久化审计的接口和状态机。等真实模型验收、审批、并发控制和审计边界都稳定后，再考虑只读广告数据适配、更多业务场景和前端审批页面。生产写入能力应当最后单独评审。")

add_heading(doc, "结束语")
add_body(doc, "总体来看，今天项目完成了从规范到可运行 PoC 的关键跨越，离线安全闭环和评估体系已经建立。当前最重要的不是继续堆功能，而是用真实模型验证现有边界是否可靠，并把审批和审计基础打牢。以上是今天的项目进展。", after=0)

doc.core_properties.title = "2026-07-17 亚马逊广告智投系统每日汇报发言稿"
doc.core_properties.subject = "亚马逊广告智投系统每日项目进展口头汇报"
doc.core_properties.author = "亚马逊广告智投系统项目组"
doc.core_properties.keywords = "Amazon Ads, PoC, 每日汇报, 2026-07-17"

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
