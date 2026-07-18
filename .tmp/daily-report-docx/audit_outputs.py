from pathlib import Path
from zipfile import ZipFile
import re

from docx import Document


ROOT = Path(r"C:\QSZ\markdown")
REPORT = ROOT / "reports" / "daily" / "2026-07-17-亚马逊广告智投系统每日报告.md"
SPEECH = ROOT / "reports" / "speech" / "2026-07-17-亚马逊广告智投系统每日汇报发言稿.docx"

expected_sections = [
    "# 亚马逊广告智投系统每日报告",
    "## 一、今日一句话总结",
    "## 二、今日完成内容",
    "## 三、项目更新明细",
    "## 四、项目推进情况",
    "## 五、今日关键成果",
    "## 六、问题、风险和阻塞项",
    "## 七、下一步计划",
    "## 八、需要重点关注的文件或模块",
    "## 九、信息来源",
]

report_text = REPORT.read_text(encoding="utf-8")
positions = [report_text.find(section) for section in expected_sections]
assert all(pos >= 0 for pos in positions), positions
assert positions == sorted(positions), positions
for fact in (
    "565 passed in 85.48s",
    "12/12 通过",
    "real_model_used=false",
    "production_write_called=false",
    "真实模型质量评估尚未执行",
    "今日未发现影响当前离线 PoC 继续开发的代码阻塞项",
):
    assert fact in report_text, fact

doc = Document(SPEECH)
paragraphs = [p for p in doc.paragraphs if p.text.strip()]
full_text = "\n".join(p.text for p in paragraphs)
spoken_headings = {"开场", "今天做了什么", "取得的关键进展", "当前项目状态", "问题与风险", "下一步", "结束语"}
opening_index = next(i for i, p in enumerate(paragraphs) if p.text == "开场")
spoken_text = "".join(p.text for p in paragraphs[opening_index + 1 :] if p.text not in spoken_headings)
chinese_chars = len(re.findall(r"[\u4e00-\u9fff]", spoken_text))

assert "五百六十五项全部通过" in full_text
assert "十二个固定离线案例全部通过" in full_text
assert "真实冒烟测试仍是 NOT EXECUTED" in full_text
assert "没有连接真实 Amazon Ads API" in full_text
assert "以上是今天的项目进展" in full_text
assert len(doc.tables) == 0
assert len(doc.sections) == 1
section = doc.sections[0]
assert round(section.page_width.inches, 2) == 8.50
assert round(section.page_height.inches, 2) == 11.00
assert all(round(value.inches, 2) == 1.00 for value in (section.top_margin, section.right_margin, section.bottom_margin, section.left_margin))
assert round(section.header_distance.inches, 3) == 0.492
assert round(section.footer_distance.inches, 3) == 0.492

with ZipFile(SPEECH) as archive:
    document_xml = archive.read("word/document.xml")
    styles_xml = archive.read("word/styles.xml")
    settings_xml = archive.read("word/settings.xml")
assert b"w:tbl" not in document_xml
assert b"Microsoft YaHei" in document_xml or b"Microsoft YaHei" in styles_xml
assert b"w:pgSz" in document_xml
assert b"w:pgMar" in document_xml

print(f"REPORT_BYTES={REPORT.stat().st_size}")
print(f"REPORT_SECTIONS={len(expected_sections)}")
print(f"DOCX_BYTES={SPEECH.stat().st_size}")
print(f"DOCX_PARAGRAPHS={len(paragraphs)}")
print(f"DOCX_TABLES={len(doc.tables)}")
print(f"SPOKEN_CHINESE_CHARS={chinese_chars}")
print("STRUCTURAL_AUDIT=PASS")
